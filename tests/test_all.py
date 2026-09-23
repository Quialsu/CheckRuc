import os
import pytest
import openpyxl
import pandas as pd
from src.validators.ruc_validator import validate_ruc, validate_ruc_checksum, process_ruc_records
from src.loaders.file_loader import load_ruc_file, extract_ruc_records_with_trace, auto_detect_ruc_column
from src.parsers.address_parser import parse_fiscal_address
from src.checkpoint.manager import CheckpointManager
from src.sources.sunat_sources import MockRUCSource, PadronReducidoSource
from src.processors.batch_processor import BatchProcessor
from src.exporters.excel_exporter import export_to_excel


def test_ruc_checksum():
    assert validate_ruc_checksum("20131312955") is True
    assert validate_ruc_checksum("20131312950") is False


def test_validate_ruc():
    is_valid, clean, err = validate_ruc("20131312955")
    assert is_valid is True
    assert clean == "20131312955"

    is_valid, clean, err = validate_ruc("")
    assert is_valid is False
    assert "vacía" in err


def test_process_ruc_records_with_trace():
    raw_recs = [
        {"archivo_origen": "test.xlsx", "fila_origen": 2, "ruc_original": "20131312955"},
        {"archivo_origen": "test.xlsx", "fila_origen": 3, "ruc_original": None},
        {"archivo_origen": "test.xlsx", "fila_origen": 4, "ruc_original": "20131312955"}
    ]
    res = process_ruc_records(raw_recs)
    assert res["total_input"] == 3
    assert res["unique_count"] == 1
    assert len(res["invalid_records"]) == 1
    assert len(res["duplicate_records"]) == 1


def test_empty_padron_source(tmp_path):
    db_file = tmp_path / "empty_padron.db"
    padron = PadronReducidoSource(db_path=str(db_file))
    res = padron.fetch_bulk_rucs(["20131312955"])
    assert res[0]["estado_consulta"] == "ERROR TEMPORAL"
    assert "no importado" in res[0]["error_tecnico"].lower()


def test_padron_versioning_and_sha256(tmp_path):
    db_file = tmp_path / "padron_versioning.db"
    padron = PadronReducidoSource(db_path=str(db_file))

    txt_file = tmp_path / "dummy_padron.txt"
    with open(txt_file, "w", encoding="latin-1") as f:
        f.write("20131312955|SUNAT TEST|ACTIVO|HABIDO|150131|AV.|CANAVAL Y MOREYRA||ZONA|150|||||\n")

    info = padron.load_from_txt_file(str(txt_file), version_name="2026-TEST-V1")
    assert len(info["sha256"]) == 64
    assert info["version"] == "2026-TEST-V1"

    res = padron.fetch_ruc("20131312955")
    assert res["razon_social"] == "SUNAT TEST"
    assert res["estado_consulta"] == "CONSULTADO"


def test_checkpoint_isolation(tmp_path):
    db_file = tmp_path / "chk_isolation.db"
    chk = CheckpointManager(db_path=str(db_file))

    chk.save_record({"ruc": "20131312955", "fuente": "MOCK", "dataset_ver": "v1.0", "razon_social": "MOCK NAME", "estado_consulta": "CONSULTADO"})
    chk.save_record({"ruc": "20131312955", "fuente": "REAL_PADRON", "dataset_ver": "2026-Q3", "razon_social": "REAL NAME", "estado_consulta": "CONSULTADO"})

    mock_recs = chk.get_processed_rucs("MOCK", "v1.0")
    real_recs = chk.get_processed_rucs("REAL_PADRON", "2026-Q3")

    assert mock_recs["20131312955"]["razon_social"] == "MOCK NAME"
    assert real_recs["20131312955"]["razon_social"] == "REAL NAME"


def test_excel_export_6_sheets(tmp_path):
    out_file = tmp_path / "Consulta_RUC_SUNAT.xlsx"
    results = [{"ruc": "20131312955", "estado_consulta": "CONSULTADO", "razon_social": "SUNAT", "fuente": "MOCK"}]
    stats = {"total_input": 2, "unique_count": 2, "consultados": 1, "no_encontrados": 0, "errores": 0, "fecha_hora": "2026-09-23", "fuente": "MOCK", "dataset_ver": "v1.0"}

    out = export_to_excel(results, [], [], stats, all_input_rucs=["20131312955", "20100000001"], output_path=str(out_file))
    wb = openpyxl.load_workbook(out)
    assert len(wb.sheetnames) == 6
    assert set(wb.sheetnames) == {"RESULTADOS", "PENDIENTES", "ERRORES", "RESUMEN", "TRAZABILIDAD", "FUENTES"}


def test_high_volume_synthetic_chunking(tmp_path):
    db_file = tmp_path / "high_vol.db"
    chk = CheckpointManager(db_path=str(db_file))

    # Generate 2,000 synthetic RUCs
    synthetic_rucs = [f"20{i:09d}" for i in range(2000)]
    stats = chk.get_summary_stats(synthetic_rucs, "MOCK", "v1.0")
    assert stats["total"] == 2000
    assert stats["pendientes"] == 2000
