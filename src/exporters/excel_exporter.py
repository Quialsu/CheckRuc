import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from typing import List, Dict, Any, Optional
from config import DEFAULT_EXCEL_OUTPUT

PRIMARY_COLUMNS = [
    "RUC",
    "Razón Social",
    "Fecha Inscripción",
    "Estado",
    "Condición",
    "Domicilio Fiscal Original",
    "Dirección",
    "Distrito",
    "Provincia",
    "Departamento/Región",
    "Ubigeo",
    "Comercio Exterior",
    "Actividad Económica Principal",
    "Actividades Económicas",
    "Fuente",
    "Fecha Consulta",
    "Estado Consulta"
]

def export_to_excel(results: List[Dict[str, Any]],
                    invalid_records: List[Dict[str, Any]],
                    duplicate_records: List[Dict[str, Any]],
                    summary_stats: Dict[str, Any],
                    all_input_rucs: Optional[List[str]] = None,
                    output_path: str = str(DEFAULT_EXCEL_OUTPUT)) -> str:
    """
    Generates the final multi-tab Excel file 'Consulta_RUC_SUNAT.xlsx' with:
      - RESULTADOS: Main records with all 17 compulsory columns formatted as text.
      - PENDIENTES: Unprocessed or pending RUCs (guaranteeing ALL pending valid RUCs are present).
      - ERRORES: Invalid RUCs, empty cells, duplicates, and technical query errors with row traceability.
      - RESUMEN: Executive overview table with total stats, timestamp, and query source version.
    """
    dir_name = os.path.dirname(output_path)
    if dir_name:
        os.makedirs(dir_name, exist_ok=True)

    wb = openpyxl.Workbook()

    # Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # -------------------------------------------------------------
    # TAB 1: RESULTADOS
    # -------------------------------------------------------------
    ws_res = wb.active
    ws_res.title = "RESULTADOS"
    ws_res.views.sheetView[0].showGridLines = True

    ws_res.append(PRIMARY_COLUMNS)
    for col_num in range(1, len(PRIMARY_COLUMNS) + 1):
        cell = ws_res.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    processed_map = {r["ruc"]: r for r in results if r.get("estado_consulta") in ["CONSULTADO", "RUC NO ENCONTRADO"]}

    row_idx = 2
    for r in processed_map.values():
        row_data = [
            str(r.get("ruc", "")),
            str(r.get("razon_social", "")),
            str(r.get("fecha_inscripcion", "")),
            str(r.get("estado", "")),
            str(r.get("condicion", "")),
            str(r.get("domicilio_fiscal_original", "")),
            str(r.get("direccion", "")),
            str(r.get("distrito", "")),
            str(r.get("provincia", "")),
            str(r.get("departamento", "") if "departamento" in r else r.get("departamento/región", "")),
            str(r.get("ubigeo", "")),
            str(r.get("comercio_exterior", "")),
            str(r.get("actividad_principal", "")),
            str(r.get("actividades_secundarias", "")),
            str(r.get("fuente", "")),
            str(r.get("fecha_consulta", "")),
            str(r.get("estado_consulta", ""))
        ]
        ws_res.append(row_data)

        ws_res.cell(row=row_idx, column=1).number_format = '@'
        for c in range(1, len(PRIMARY_COLUMNS) + 1):
            cell = ws_res.cell(row=row_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
        row_idx += 1

    ws_res.freeze_panes = "A2"
    ws_res.auto_filter.ref = f"A1:{get_column_letter(len(PRIMARY_COLUMNS))}{max(len(processed_map)+1, 1)}"

    # -------------------------------------------------------------
    # TAB 2: PENDIENTES
    # -------------------------------------------------------------
    ws_pen = wb.create_sheet(title="PENDIENTES")
    ws_pen.views.sheetView[0].showGridLines = True
    pen_cols = ["RUC", "Estado Consulta", "Detalle"]
    ws_pen.append(pen_cols)

    for c in range(1, len(pen_cols) + 1):
        cell = ws_pen.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    all_input_rucs = all_input_rucs or []
    pending_rucs = [r for r in all_input_rucs if r not in processed_map]

    pen_idx = 2
    for ruc in pending_rucs:
        ws_pen.append([str(ruc), "PENDIENTE", "Pendiente de procesamiento / interrupción"])
        ws_pen.cell(row=pen_idx, column=1).number_format = '@'
        for c in range(1, len(pen_cols) + 1):
            cell = ws_pen.cell(row=pen_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
        pen_idx += 1

    ws_pen.freeze_panes = "A2"

    # -------------------------------------------------------------
    # TAB 3: ERRORES
    # -------------------------------------------------------------
    ws_err = wb.create_sheet(title="ERRORES")
    ws_err.views.sheetView[0].showGridLines = True
    err_cols = ["Tipo Error", "Archivo", "Fila Origen", "Valor Original", "Motivo Técnico"]
    ws_err.append(err_cols)
    for c in range(1, len(err_cols) + 1):
        cell = ws_err.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    err_row_idx = 2
    for inv in invalid_records:
        ws_err.append(["RUC Inválido / Celda Vacía", str(inv.get('archivo', '-')), f"Fila {inv.get('fila', '-')}", str(inv.get('ruc_original', '')), str(inv.get('motivo', ''))])
        for c in range(1, len(err_cols) + 1):
            cell = ws_err.cell(row=err_row_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
        err_row_idx += 1

    for dup in duplicate_records:
        ws_err.append(["Duplicado Omite", str(dup.get('archivo', '-')), f"Fila {dup.get('fila', '-')}", str(dup.get('ruc_original', '')), str(dup.get('motivo', ''))])
        for c in range(1, len(err_cols) + 1):
            cell = ws_err.cell(row=err_row_idx, column=c)
            cell.font = data_font
            cell.border = thin_border
        err_row_idx += 1

    for r in results:
        if r.get("estado_consulta") in ["ERROR TEMPORAL", "REQUIERE REVISIÓN"]:
            ws_err.append(["Error Consulta", "-", "-", str(r.get('ruc', '')), str(r.get('error_tecnico', ''))])
            for c in range(1, len(err_cols) + 1):
                cell = ws_err.cell(row=err_row_idx, column=c)
                cell.font = data_font
                cell.border = thin_border
            err_row_idx += 1

    ws_err.freeze_panes = "A2"

    # -------------------------------------------------------------
    # TAB 4: RESUMEN
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="RESUMEN")
    ws_sum.views.sheetView[0].showGridLines = True
    ws_sum.append(["INDICADOR DE EJECUCIÓN", "VALOR / DETALLE"])

    for c in range(1, 3):
        cell = ws_sum.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font

    summary_rows = [
        ("Total RUCs Ingresados", summary_stats.get("total_input", 0)),
        ("RUCs Únicos Válidos", summary_stats.get("unique_count", 0)),
        ("RUCs Consultados Exitosos", summary_stats.get("consultados", 0)),
        ("RUCs No Encontrados", summary_stats.get("no_encontrados", 0)),
        ("RUCs con Error / Requieren Revisión", summary_stats.get("errores", 0)),
        ("RUCs Pendientes", len(pending_rucs)),
        ("Fecha / Hora Procesamiento", summary_stats.get("fecha_hora", "")),
        ("Fuente Utilizada", summary_stats.get("fuente", "")),
        ("Versión de Dataset", summary_stats.get("dataset_ver", ""))
    ]

    for row_idx, (k, v) in enumerate(summary_rows, start=2):
        ws_sum.append([k, v])
        cell_k = ws_sum.cell(row=row_idx, column=1)
        cell_v = ws_sum.cell(row=row_idx, column=2)
        cell_k.font = Font(name="Calibri", size=10, bold=True)
        cell_v.font = data_font
        cell_k.border = thin_border
        cell_v.border = thin_border

    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(output_path)
    return output_path
